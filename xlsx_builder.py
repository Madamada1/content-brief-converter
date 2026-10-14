"""
Builds a Myer-format content-brief workbook (.xlsx) from the structured
dict produced by docx_parser.parse_docx().

Sheet 1 mirrors Myer's real "SEO CONTENT PAGE" template exactly (the same
layout used in the SS27 workbook's "SEO TEMPLATE" tab and filled examples
like "WK06 TOP BLENDERS 2026" / "WK06 BEST IRONS"):
  - a header block (GO-LIVE DATE, JAYWING BRIEF, RESKIN OR NEW, URL, META
    TITLE, META DESCRIPTION, ANCHOR HEADINGS)
  - a MODULE table with a "COPY HERO ROW" for the H1, then one
    " CONTENT ROW N" block per H2/H3 section, a "DY STRIP" block
    immediately under any section that calls for a product carousel, and
    a "FAQs" block for the FAQ questions/answers.

Sheet 2 "Keywords & Notes" keeps the reference info from the brief
(job purpose, target audience, keyword research, etc.) that isn't part of
the template but is useful to keep alongside it.
"""

import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    RICH_TEXT_AVAILABLE = True
except ImportError:  # openpyxl < 3.1 doesn't have rich text support
    RICH_TEXT_AVAILABLE = False

# --- Font ---------------------------------------------------------------

FONT_NAME = "Poppins"

# --- Shared styling ---------------------------------------------------

LABEL_FONT = Font(name=FONT_NAME, bold=True)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
COLHEADER_FONT = Font(name=FONT_NAME, bold=True)
COLHEADER_FILL = PatternFill("solid", fgColor="D9E2E8")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
PLAIN_FONT = Font(name=FONT_NAME, size=11)

# --- Content-brief sheet styling (matches the real SEO CONTENT PAGE tabs) --

TITLE_FILL = PatternFill("solid", fgColor="5A5A5A")
TITLE_TEXT_FONT = Font(name=FONT_NAME, bold=True, size=15, color="FFFFFFFF")
FIELD_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=12)
FIELD_VALUE_FONT = Font(name=FONT_NAME, size=12)
FIELD_VALUE_ALIGN = Alignment(vertical="top", wrap_text=True)

MODULE_HEADER_FILL = PatternFill("solid", fgColor="D7E2DC")
MODULE_HEADER_FONT = Font(name=FONT_NAME, bold=True, size=14, color="FF000000")
MODULE_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

DIVIDER_FILL = PatternFill("solid", fgColor="B4CDC9")
DIVIDER_FONT = Font(name=FONT_NAME, bold=True, size=13, color="FF000000")

DY_FILL = PatternFill("solid", fgColor="C8DFE5")
DY_FONT = Font(name=FONT_NAME, size=12, color="FF000000")

HEADING_CELL_FONT = Font(name=FONT_NAME, bold=True, size=12, color="FF000000")
BODY_CELL_FONT = Font(name=FONT_NAME, size=12, color="FF000000")
BODY_LINE_HEIGHT = 19

# Brand-logo / category link-grid rows (e.g. "BRAND LOGO ROW", "CATEGORY ROW")
ROW_HEADLINE_FONT = Font(name=FONT_NAME, bold=True, size=13, color="FFC90CBD")
LINK_FILL = PatternFill("solid", fgColor="EFE8FF")
LINK_FONT = Font(name=FONT_NAME, size=12, color="FF0563C1", underline="single")

CONTENT_COL_WIDTHS = [27, 32.71, 87.86, 122.29, 30, 40, 60, 3.57, 70]
CONTENT_COL_CHARS = {3: 88, 4: 122, 6: 40, 7: 60, 9: 70}  # for row-height estimation

COL_CARRYOVER, COL_MODULE, COL_HEADLINE, COL_SUBCOPY = 1, 2, 3, 4
COL_ASSET, COL_CAMPAIGN, COL_NOTES, COL_SPACER, COL_LINK = 5, 6, 7, 8, 9

PRODUCT_CAROUSEL_RE = re.compile(r"product\s+carousel.*?grid to (.+?)\]?\s*$", re.IGNORECASE)
IMAGE_GRID_RE = re.compile(r"internal linking image grid", re.IGNORECASE)
URL_RE = re.compile(r"https?://[^\s\]]+")


def _inline_font(bold=False, underline=False):
    return InlineFont(rFont=FONT_NAME, b=bold, u=("single" if underline else None))


def _build_rich_copy(rich_paragraphs):
    """Turn docx_parser's [{"runs": [TextRun,...], "bullet": bool}, ...]
    into an openpyxl CellRichText: Word bold -> bold run, Word bullet -> a
    literal '• ' prefix on its own line, a blank line between a bullet list
    and the body copy either side of it, Word hyperlink -> underlined black
    run (no hyperlink-blue colour, no visible URL). Returns None if rich
    text isn't available/usable, so the caller can fall back to a plain
    string.

    IMPORTANT: openpyxl only marks a <t> run as xml:space="preserve" when
    its *stripped* text is non-empty (see openpyxl.xml.functions.
    whitespace) -- a run whose text is pure whitespace (e.g. a bare "\\n"
    separator) never gets that flag, and Excel then silently discards the
    run's text entirely, which is what was collapsing paragraph/bullet
    boundaries together. To avoid ever emitting a whitespace-only run, every
    newline separator is merged into an adjacent run that has real
    (non-whitespace) content -- the bullet glyph for bullet paragraphs, or
    the first real word for non-bullet paragraphs.
    """
    if not RICH_TEXT_AVAILABLE or not rich_paragraphs:
        return None

    # Drop blank/whitespace-only paragraphs and runs (soft-break noise) so
    # they can't end up as unanchored whitespace-only blocks either.
    paragraphs = []
    for para in rich_paragraphs:
        runs = [r for r in para["runs"] if r.text and r.text.strip()]
        if runs:
            paragraphs.append({"runs": runs, "bullet": bool(para.get("bullet"))})
    if not paragraphs:
        return None

    blocks = []
    prev_bullet = None
    for i, para in enumerate(paragraphs):
        separator = ""
        if i > 0:
            separator = "\n\n" if para["bullet"] != prev_bullet else "\n"

        runs = para["runs"]
        if para["bullet"]:
            # The bullet glyph is real (non-whitespace) content, so merging
            # the separator into it is always safe.
            blocks.append(TextBlock(_inline_font(), separator + "• "))
        elif separator:
            # No bullet glyph to anchor to -- merge into the first run's
            # own text instead, keeping that run's own bold/underline.
            first, runs = runs[0], runs[1:]
            blocks.append(
                TextBlock(_inline_font(bold=first.bold, underline=first.underline), separator + first.text)
            )

        for run in runs:
            blocks.append(TextBlock(_inline_font(bold=run.bold, underline=run.underline), run.text))

        prev_bullet = para["bullet"]

    return CellRichText(blocks) if blocks else None


def _estimate_row_height(text, col_width_chars, line_height=15):
    if not text:
        return 15
    lines = 0
    for line in str(text).split("\n"):
        lines += max(1, math.ceil(len(line) / max(col_width_chars, 10)))
    return min(max(15, lines * line_height), 409)  # 409 is Excel's row-height cap


def _set_col_widths(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


# --- Generic label-row helper, used by the Keywords & Notes sheet ---------


def _label_row(ws, row, label, value="", note=None):
    ws.cell(row=row, column=1, value=label).font = LABEL_FONT
    cell_b = ws.cell(row=row, column=2, value=value)
    cell_b.font = PLAIN_FONT
    cell_b.alignment = WRAP_TOP
    if note:
        cell_c = ws.cell(row=row, column=3, value=note)
        cell_c.font = PLAIN_FONT
        cell_c.alignment = WRAP_TOP
    ws.row_dimensions[row].height = _estimate_row_height(value, 70)
    return row + 1


def build_workbook(parsed):
    wb = Workbook()
    _build_content_brief_sheet(wb.active, parsed)
    _build_keywords_sheet(wb.create_sheet("Keywords & Notes"), parsed.get("brief_info", {}))
    return wb


# --- Content-brief sheet ---------------------------------------------


def _build_content_brief_sheet(ws, parsed):
    ws.title = _safe_sheet_title(parsed.get("content_page_name"))
    _set_col_widths(ws, CONTENT_COL_WIDTHS)

    row = _write_title_row(ws, 1)
    row = _write_field_row(ws, row, "GO-LIVE DATE", "", height=28.5)
    source_file = parsed.get("source_filename") or ""
    row = _write_field_row(ws, row, "JAYWING BRIEF", source_file, height=28.5)
    row = _write_field_row(ws, row, "RESKIN OR NEW", "", height=28.5)
    row += 1  # blank gap row, matching the real template

    info = parsed.get("brief_info", {})
    row = _write_field_row(ws, row, "URL", info.get("existing_url", ""), height=28.5)
    row = _write_field_row(ws, row, "META TITLE", info.get("meta_title", ""), height=28.5)
    row = _write_field_row(
        ws, row, "META DESCRIPTION", info.get("meta_description", ""), height=30
    )
    row = _write_anchor_row(ws, row, parsed.get("anchors") or [])

    row = _write_module_header_row(ws, row)
    row = _write_content_blocks(ws, row, parsed.get("sections", []))

    ws.freeze_panes = "A11"


def _safe_sheet_title(name, fallback="Content Brief"):
    name = (name or fallback).strip() or fallback
    name = re.sub(r"[\\/*?:\[\]]", "", name)
    return name[:31] or fallback


def _write_title_row(ws, row):
    cell = ws.cell(row=row, column=1, value="SEO CONTENT PAGE")
    cell.font = TITLE_TEXT_FONT
    for col in range(1, len(CONTENT_COL_WIDTHS) + 1):
        ws.cell(row=row, column=col).fill = TITLE_FILL
    ws.row_dimensions[row].height = 38.25
    return row + 1


def _write_field_row(ws, row, label, value, height=28.5):
    label_cell = ws.cell(row=row, column=1, value=label)
    label_cell.font = FIELD_LABEL_FONT
    value_cell = ws.cell(row=row, column=2, value=value)
    value_cell.font = FIELD_VALUE_FONT
    value_cell.alignment = FIELD_VALUE_ALIGN
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(CONTENT_COL_WIDTHS))
    ws.row_dimensions[row].height = max(height, _estimate_row_height(value, 150))
    return row + 1


def _write_anchor_row(ws, row, anchors):
    label_cell = ws.cell(row=row, column=1, value="ANCHOR HEADINGS")
    label_cell.font = FIELD_LABEL_FONT
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
    value_cell = ws.cell(row=row, column=4, value=" | ".join(anchors))
    value_cell.font = FIELD_VALUE_FONT
    value_cell.alignment = FIELD_VALUE_ALIGN
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=len(CONTENT_COL_WIDTHS))
    ws.row_dimensions[row].height = 30
    return row + 1


def _write_module_header_row(ws, row):
    headers = {
        COL_CARRYOVER: "CARRYOVER OR NEW",
        COL_MODULE: "MODULE",
        COL_HEADLINE: "HEADLINE COPY DIRECTION",
        COL_SUBCOPY: "SUBCOPY COPY DIRECTION",
        COL_ASSET: "ASSET ID(S)",
        COL_CAMPAIGN: "CAMPAIGN DESIGN TREATMENT/CTA",
        COL_NOTES: "ADDITIONAL CREATIVE DESIGN NOTES",
        COL_LINK: "LINK",
    }
    for col, text in headers.items():
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = MODULE_HEADER_FONT
        cell.alignment = MODULE_HEADER_ALIGN
    for col in range(1, len(CONTENT_COL_WIDTHS) + 1):
        ws.cell(row=row, column=col).fill = MODULE_HEADER_FILL
    ws.row_dimensions[row].height = 42
    return row + 1


def _divider_row(ws, row, text_a, text_b=None):
    cell_a = ws.cell(row=row, column=1, value=text_a)
    cell_a.font = DIVIDER_FONT
    if text_b is not None:
        cell_b = ws.cell(row=row, column=2, value=text_b)
        cell_b.font = DIVIDER_FONT
    for col in range(1, len(CONTENT_COL_WIDTHS) + 1):
        ws.cell(row=row, column=col).fill = DIVIDER_FILL
    ws.row_dimensions[row].height = 33.75
    return row + 1


def _content_data_row(ws, row, values, bold_cols=(), rich_cols=None):
    """Write a MODULE data row. `values` maps column -> plain text (used for
    every column, and as the row-height / fallback source for rich
    columns). `rich_cols` optionally maps column -> docx_parser rich
    paragraphs (see _build_rich_copy) for columns that should render Word's
    bold/bullet/hyperlink formatting instead of the plain string."""
    rich_cols = rich_cols or {}
    ws.cell(row=row, column=COL_CARRYOVER, value="NEW CONTENT").font = BODY_CELL_FONT
    max_height = 24
    for col, value in values.items():
        cell = ws.cell(row=row, column=col)
        rich = _build_rich_copy(rich_cols.get(col)) if col in rich_cols else None
        if rich is not None:
            cell.value = rich
        else:
            cell.value = value
        cell.font = HEADING_CELL_FONT if col in bold_cols else BODY_CELL_FONT
        cell.alignment = FIELD_VALUE_ALIGN
        chars = CONTENT_COL_CHARS.get(col, 40)
        # Size-12 Poppins wraps at ~19pt per line; 15pt clipped the last lines.
        max_height = max(max_height, _estimate_row_height(value, chars, line_height=BODY_LINE_HEIGHT))
    ws.row_dimensions[row].height = max_height
    return row + 1


def _dy_strip_block(ws, row, note_text):
    match = PRODUCT_CAROUSEL_RE.search(note_text)
    title = match.group(1).strip(" .").title() if match else "Product Picks"
    row = _divider_row(ws, row, "DY STRIP", title)
    ws.cell(row=row, column=COL_CARRYOVER, value="NEW CONTENT").font = DY_FONT
    ws.cell(row=row, column=COL_MODULE, value="DY Strip").font = DY_FONT
    ws.cell(row=row, column=COL_HEADLINE, value="DY CODE: #tbc").font = DY_FONT
    for col in range(1, len(CONTENT_COL_WIDTHS) + 1):
        ws.cell(row=row, column=col).fill = DY_FILL
    ws.row_dimensions[row].height = 30
    return row + 1


def _product_grid_block(ws, row, products):
    """Renders docx_parser's detected product-recommendation entries
    (a repeating URL + bold name + description pattern in the source
    copy) as a PRODUCT GRID block: one divider, then one MODULE row per
    product with its name, description, and a real clickable link in the
    LINK column -- instead of leaving all the products flattened into a
    single SUBCOPY cell."""
    row = _divider_row(ws, row, "PRODUCT GRID")
    for product in products:
        written_row = row
        row = _content_data_row(
            ws,
            row,
            {
                COL_MODULE: "6UP Scroll",
                COL_HEADLINE: product["name"] or "NA",
                COL_SUBCOPY: product["copy"],
                COL_LINK: product["url"] or "",
            },
            rich_cols={COL_SUBCOPY: product.get("copy_runs")},
        )
        if product["url"]:
            link_cell = ws.cell(row=written_row, column=COL_LINK)
            link_cell.hyperlink = product["url"]
            link_cell.font = LINK_FONT
            link_cell.fill = LINK_FILL
    return row


def _products_and_outro(ws, row, section):
    """Write a section's PRODUCT GRID (if any), then any body copy that
    followed the last product in the brief as a continuation Copy row."""
    if section.get("products"):
        row = _product_grid_block(ws, row, section["products"])
    if section.get("copy_after", "").strip():
        row = _content_data_row(
            ws,
            row,
            {COL_MODULE: "Copy", COL_HEADLINE: "NA", COL_SUBCOPY: section["copy_after"], COL_ASSET: "NA"},
            rich_cols={COL_SUBCOPY: section.get("copy_after_runs")},
        )
    return row


def _url_to_label(url):
    """Turn a category/brand URL into a readable tile name, e.g.
    '.../bedding/quilt-covers' -> 'Quilt Covers'."""
    path = url.rstrip("/").split("?")[0]
    slug = path.rstrip("/").split("/")[-1]
    words = re.split(r"[-_]+", slug)
    label = " ".join(w.capitalize() for w in words if w)
    return label or url


def _image_grid_block(ws, row, title, note_text):
    """Renders a "[Myer to Insert internal linking image grid ...]" note as
    a proper BRAND LOGO ROW / CATEGORY ROW block (matching the real
    template) instead of dumping the instruction text into a notes column.
    Brand-grid notes have no named brands to link to, so that block is left
    as a scaffold (divider + headline, no fabricated tiles). Category-grid
    notes in this brief template list real example URLs inline, so those
    become real clickable link tiles."""
    is_brand = "brand" in note_text.lower()
    module = "6UP Scroll" if is_brand else "4UP Scroll"
    urls = URL_RE.findall(note_text)

    row = _divider_row(ws, row, "BRAND LOGO ROW" if is_brand else "CATEGORY ROW")

    ws.cell(row=row, column=1, value="Row Headline H2:").font = FIELD_LABEL_FONT
    headline_cell = ws.cell(row=row, column=2, value=title)
    headline_cell.font = ROW_HEADLINE_FONT
    ws.row_dimensions[row].height = 24
    row += 1

    for url in urls:
        ws.cell(row=row, column=COL_CARRYOVER, value="NEW CONTENT").font = BODY_CELL_FONT
        ws.cell(row=row, column=COL_MODULE, value=module).font = BODY_CELL_FONT
        ws.cell(row=row, column=COL_HEADLINE, value=_url_to_label(url)).font = BODY_CELL_FONT
        link_cell = ws.cell(row=row, column=COL_LINK, value=url)
        link_cell.hyperlink = url
        link_cell.font = LINK_FONT
        link_cell.fill = LINK_FILL
        ws.row_dimensions[row].height = 24
        row += 1

    return row


def _write_content_blocks(ws, row, sections):
    content_row_num = 0
    i = 0
    n = len(sections)

    while i < n:
        section = sections[i]

        if section["level"] == 1:
            row = _divider_row(ws, row, " COPY HERO ROW")
            row = _content_data_row(
                ws,
                row,
                {
                    COL_MODULE: "H1 Heading",
                    COL_HEADLINE: section["heading"],
                    COL_SUBCOPY: section["copy"],
                    COL_ASSET: "NA",
                },
                rich_cols={COL_SUBCOPY: section.get("copy_runs")},
            )
            row = _products_and_outro(ws, row, section)
            i += 1
            continue

        if section["level"] == 2 and section["heading"].strip().lower() == "faqs":
            row = _divider_row(ws, row, "FAQs")
            row = _divider_row(ws, row, "Row Headline H1:", "FAQs")
            i += 1
            while i < n and sections[i]["level"] == 3:
                faq = sections[i]
                row = _content_data_row(
                    ws,
                    row,
                    {
                        COL_MODULE: "Copy",
                        COL_HEADLINE: faq["heading"],
                        COL_SUBCOPY: faq["copy"],
                        COL_ASSET: "NA",
                    },
                    rich_cols={COL_SUBCOPY: faq.get("copy_runs")},
                )
                row = _products_and_outro(ws, row, faq)
                i += 1
            continue

        # Standard H2/H3 content row
        note = section.get("note") or ""
        is_carousel = bool(PRODUCT_CAROUSEL_RE.search(note))
        is_image_grid = bool(IMAGE_GRID_RE.search(note))

        if is_image_grid:
            # "[Myer to Insert internal linking image grid ...]" -> a real
            # BRAND LOGO ROW / CATEGORY ROW block, not a dumped note. If the
            # section also has its own intro copy, show that as a normal
            # content row first.
            if section["copy"].strip():
                content_row_num += 1
                row = _divider_row(ws, row, f" CONTENT ROW {content_row_num}")
                row = _content_data_row(
                    ws,
                    row,
                    {
                        COL_MODULE: "Copy",
                        COL_HEADLINE: f"H{section['level']} HEADING: {section['heading']}",
                        COL_SUBCOPY: section["copy"],
                        COL_ASSET: "NA",
                    },
                    bold_cols={COL_HEADLINE},
                    rich_cols={COL_SUBCOPY: section.get("copy_runs")},
                )
            row = _image_grid_block(ws, row, section["heading"], note)
            i += 1
            continue

        content_row_num += 1
        row = _divider_row(ws, row, f" CONTENT ROW {content_row_num}")
        row = _content_data_row(
            ws,
            row,
            {
                COL_MODULE: "Copy",
                COL_HEADLINE: f"H{section['level']} HEADING: {section['heading']}",
                COL_SUBCOPY: section["copy"],
                COL_ASSET: "NA",
                COL_NOTES: "" if is_carousel else note,
            },
            bold_cols={COL_HEADLINE},
            rich_cols={COL_SUBCOPY: section.get("copy_runs")},
        )
        if is_carousel:
            row = _dy_strip_block(ws, row, note)
        row = _products_and_outro(ws, row, section)
        i += 1

    return row


# --- Keywords & Notes sheet --------------------------------------------


def _build_keywords_sheet(ws, info):
    _set_col_widths(ws, [26, 55, 20, 20])

    row = 1
    ws.cell(row=row, column=1, value="Brief Reference Info").font = TITLE_FONT
    row += 2

    row = _label_row(ws, row, "Copywriter", info.get("copywriter", ""))
    row = _label_row(ws, row, "Client", info.get("client", ""))
    row = _label_row(ws, row, "Date", info.get("date", ""))
    row += 1

    row = _label_row(ws, row, "Job Purpose", info.get("job_purpose", ""))
    row = _label_row(ws, row, "Type of Content", info.get("type_of_content", ""))
    row = _label_row(ws, row, "Brand & TOV Guidelines", "\n".join(info.get("brand_tov", [])))
    row = _label_row(ws, row, "Target Audience", info.get("target_audience", ""))
    row += 1

    row = _label_row(
        ws,
        row,
        "Meta Title",
        info.get("meta_title", ""),
        note=info.get("meta_title_pixels") or None,
    )
    row = _label_row(
        ws,
        row,
        "Meta Description",
        info.get("meta_description", ""),
        note=info.get("meta_description_pixels") or None,
    )
    row += 1

    keywords = info.get("keywords") or []
    if keywords:
        header_row = row
        headers = ["Type", "Keyword", "Avg. Monthly Searches", "Ranking"]
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=header_row, column=col, value=h)
            c.font = COLHEADER_FONT
            c.fill = COLHEADER_FILL
        row += 1
        for kw_row in keywords:
            for col, value in enumerate(kw_row, start=1):
                ws.cell(row=row, column=col, value=value).font = PLAIN_FONT
            row += 1
        row += 1

    related = info.get("related_keywords") or []
    if related:
        row = _label_row(ws, row, "Semantically Related Keywords", ", ".join(related))


def safe_filename(name, fallback="content-brief"):
    name = (name or fallback).strip() or fallback
    name = re.sub(r"[^\w\- ]+", "", name)
    name = re.sub(r"\s+", "-", name).strip("-")
    return (name or fallback).lower()

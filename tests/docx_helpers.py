"""Helpers for building minimal .docx fixtures in tests, mimicking the
Jaywing brief template's "Draft Content" section (see docx_parser.py)."""

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn


def add_hyperlink(paragraph, url, text, bold=False):
    """Append a real Word hyperlink run (not just styled text) to `paragraph`."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    if bold:
        r_pr.append(OxmlElement("w:b"))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_bullet(paragraph):
    """Mark `paragraph` as a Word bulleted-list item (adds a <w:numPr>)."""
    numpr_xml = (
        '<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
    )
    paragraph._p.get_or_add_pPr().append(parse_xml(numpr_xml))


def build_brief_docx(path, heading, intro_lines, products, bullet_url=False, bold_name=False,
                      trailing_heading=None, trailing_lines=None):
    """Build a minimal Jaywing-style brief .docx with a single H2 section
    containing `intro_lines` of body copy followed by a repeating product
    grid built from `products` (list of (url, name, [detail_lines...])).

    `bullet_url` / `bold_name` control whether the URL paragraph is a
    bulleted list item and whether the name paragraph is bold -- both
    should be optional, since docx_parser must detect the product-grid
    pattern from the Myer product URL regardless of that formatting.
    """
    doc = Document()
    doc.add_paragraph("Draft Content")
    doc.add_paragraph(f"H2: {heading}")
    for line in intro_lines:
        doc.add_paragraph(line)

    for url, name, detail_lines in products:
        p = doc.add_paragraph()
        if bullet_url:
            make_bullet(p)
        add_hyperlink(p, url, url)

        name_p = doc.add_paragraph()
        if bold_name:
            run = name_p.add_run(name)
            run.bold = True
        else:
            name_p.add_run(name)

        for line in detail_lines:
            doc.add_paragraph(line)

    if trailing_heading:
        doc.add_paragraph(f"H2: {trailing_heading}")
        for line in trailing_lines or []:
            doc.add_paragraph(line)

    doc.save(path)
    return path

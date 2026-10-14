"""Tests for product grids written as repeated labelled fields
("Name: / Type: / Pros: / Cons: / Star Rating: / Price:"), with the
product link hyperlinked on the name rather than on its own line.

See docx_parser._split_field_records and docx_parser.clean_url.
"""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx_helpers import add_hyperlink, make_bullet

from docx_parser import clean_url, parse_docx
from xlsx_builder import COL_HEADLINE, COL_LINK, COL_SUBCOPY, build_workbook

MACHINES = [
    (
        "Nespresso by Breville Vertuo Creatista",
        # Mangled Word Online link, as found in real Jaywing briefs.
        "https://auc-word-edit.officeapps.live.com/we/8.%09https:/www.myer.com.au/p/nespresso-vertuo-creatista",
        "https://www.myer.com.au/p/nespresso-vertuo-creatista",
        ["Type: Capsule", "Pros: Easy milk drinks", "Cons: Vertuo capsules only",
         "Star Rating: 4.5/5", "Price: $879.20"],
    ),
    (
        "KitchenAid KF4",
        "https://www.myer.com.au/p/kitchenaid-kf4",
        "https://www.myer.com.au/p/kitchenaid-kf4",
        ["Type: Fully automatic", "Pros: Built-in grinder", "Cons: Milk system needs cleaning",
         "Star Rating: 5/5", "Price: $1,599"],
    ),
    (
        "Philips Baristina",
        "https://www.myer.com.au/p/philips-baristina",
        "https://www.myer.com.au/p/philips-baristina",
        ["Type: Semi-automatic", "Pros: Compact and quiet", "Cons: Limited manual control",
         "Star Rating: 4.3/5", "Price: $599"],
    ),
]


def _bullet(doc, text):
    p = doc.add_paragraph(text)
    make_bullet(p)
    return p


def build_field_grid_docx(path, machines=MACHINES, labelled_name=True, outro_lines=()):
    doc = Document()
    doc.add_paragraph("Draft Content")
    doc.add_paragraph("H2: The Best Coffee Machines to Buy")
    doc.add_paragraph("Here's how some of the best machines compare.")
    for name, raw_url, _, fields in machines:
        p = doc.add_paragraph()
        make_bullet(p)
        if labelled_name:
            p.add_run("Name: ")
        add_hyperlink(p, raw_url, name)
        for line in fields:
            _bullet(doc, line)
    for line in outro_lines:
        doc.add_paragraph(line)
    doc.add_paragraph("H2: How to Choose")
    doc.add_paragraph("The right machine should suit your routine.")
    for line in ["Budget: Allow for ongoing costs.", "Size: Measure your bench space.",
                 "Grinder: Built-in saves space."]:
        _bullet(doc, line)
    doc.save(path)
    return path


class CleanUrlTest(unittest.TestCase):
    def test_word_online_wrapper_removed(self):
        self.assertEqual(
            clean_url("https://auc-word-edit.officeapps.live.com/we/-%09https:/www.myer.com.au/c/home/shavers"),
            "https://www.myer.com.au/c/home/shavers",
        )

    def test_normal_url_unchanged(self):
        self.assertEqual(clean_url("https://www.myer.com.au/p/abc"), "https://www.myer.com.au/p/abc")


class FieldRecordGridTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)

    def _parse(self, **kwargs):
        path = os.path.join(self.tmpdir.name, "coffee.docx")
        build_field_grid_docx(path, **kwargs)
        return parse_docx(path)

    def test_name_field_records_become_products_with_hyperlink_urls(self):
        grid, choose = self._parse()["sections"]
        self.assertEqual(grid["copy"].strip(), "Here's how some of the best machines compare.")
        self.assertEqual(len(grid["products"]), len(MACHINES))
        for (name, _, url, fields), product in zip(MACHINES, grid["products"]):
            self.assertEqual(product["name"], name)
            self.assertEqual(product["url"], url)
            self.assertNotIn("Name:", product["copy"])
            for line in fields:
                self.assertIn(line, product["copy"])

    def test_one_off_labelled_tips_are_not_a_grid(self):
        _, choose = self._parse()["sections"]
        self.assertEqual(choose["products"], [])
        self.assertIn("Budget: Allow for ongoing costs.", choose["copy"])

    def test_closing_copy_after_grid_is_kept_separately(self):
        grid, _ = self._parse(outro_lines=["Still unsure? Visit a Myer store."])["sections"]
        self.assertEqual(len(grid["products"]), len(MACHINES))
        self.assertNotIn("Still unsure", grid["products"][-1]["copy"])
        self.assertIn("Still unsure? Visit a Myer store.", grid["copy_after"])

    def test_unlabelled_hyperlinked_name_above_fields(self):
        grid, _ = self._parse(labelled_name=False)["sections"]
        self.assertEqual([p["name"] for p in grid["products"]], [m[0] for m in MACHINES])
        self.assertEqual([p["url"] for p in grid["products"]], [m[2] for m in MACHINES])
        self.assertTrue(grid["copy"].startswith("Here's how"))

    def test_excel_rows_have_name_copy_and_link(self):
        path = os.path.join(self.tmpdir.name, "coffee.docx")
        build_field_grid_docx(path, outro_lines=["Still unsure? Visit a Myer store."])
        parsed = parse_docx(path)
        parsed["source_filename"] = "coffee.docx"
        ws = build_workbook(parsed).active

        divider = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "PRODUCT GRID")
        for offset, (name, _, url, fields) in enumerate(MACHINES, start=1):
            row = divider + offset
            self.assertEqual(ws.cell(row=row, column=COL_HEADLINE).value, name)
            self.assertEqual(ws.cell(row=row, column=COL_LINK).value, url)
            self.assertEqual(ws.cell(row=row, column=COL_LINK).hyperlink.target, url)
            for line in fields:
                self.assertIn(line, str(ws.cell(row=row, column=COL_SUBCOPY).value))
        outro_row = divider + len(MACHINES) + 1
        self.assertIn("Still unsure", str(ws.cell(row=outro_row, column=COL_SUBCOPY).value))


if __name__ == "__main__":
    unittest.main()
